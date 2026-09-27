import abc
import uuid
import pypdf
import docx
from typing import Dict, Type
from pathlib import Path

from app.models.schemas import NormalizedDocument, NormalizedPage, Block

class BaseParser(abc.ABC):
    @abc.abstractmethod
    def parse(self, file_path: str, document_id: str) -> NormalizedDocument:
        pass

class TxtParser(BaseParser):
    def parse(self, file_path: str, document_id: str) -> NormalizedDocument:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
            
        block = Block(
            block_id=str(uuid.uuid4()),
            block_type="paragraph",
            content=text,
            page=1
        )
        
        page = NormalizedPage(page_number=1, blocks=[block])
        
        return NormalizedDocument(
            document_id=document_id,
            pages=[page],
            metadata={"source_type": "txt"}
        )

class PdfParser(BaseParser):
    def parse(self, file_path: str, document_id: str) -> NormalizedDocument:
        doc = NormalizedDocument(document_id=document_id, metadata={"source_type": "pdf"})
        
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            doc.metadata["total_pages"] = len(reader.pages)
            
            for i, pdf_page in enumerate(reader.pages):
                page_num = i + 1
                text = pdf_page.extract_text() or ""
                
                # Basic chunking by paragraph (double newline)
                paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
                
                blocks = []
                for p in paragraphs:
                    blocks.append(Block(
                        block_id=str(uuid.uuid4()),
                        block_type="paragraph",
                        content=p,
                        page=page_num
                    ))
                
                doc.pages.append(NormalizedPage(page_number=page_num, blocks=blocks))
                
        return doc

class DocxParser(BaseParser):
    def parse(self, file_path: str, document_id: str) -> NormalizedDocument:
        doc = NormalizedDocument(document_id=document_id, metadata={"source_type": "docx"})
        word_doc = docx.Document(file_path)
        
        blocks = []
        for i, para in enumerate(word_doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue
                
            block_type = "paragraph"
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                block_type = "heading"
                
            blocks.append(Block(
                block_id=str(uuid.uuid4()),
                block_type=block_type,
                content=text,
                page=1, # DOCX doesn't easily expose pages
                metadata={"style": para.style.name if para.style else None}
            ))

        for i, table in enumerate(word_doc.tables):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(" | ".join(row_data))
            
            table_text = "\n".join(table_data).strip()
            if table_text:
                blocks.append(Block(
                    block_id=str(uuid.uuid4()),
                    block_type="table",
                    content=table_text,
                    page=1,
                    metadata={"table_index": i}
                ))
                
        for rel in word_doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    image_part = rel.target_part
                    asset_id = str(uuid.uuid4())
                    ext = Path(image_part.partname).suffix if image_part.partname else ".png"
                    filename = f"{asset_id}{ext}"
                    
                    doc.assets.append({
                        "asset_id": asset_id,
                        "asset_type": "image",
                        "mime_type": getattr(image_part, "content_type", "image/png"),
                        "filename": filename,
                        "data": image_part.blob
                    })
                except Exception as e:
                    # Ignore failed image extractions gracefully
                    pass

        doc.pages.append(NormalizedPage(page_number=1, blocks=blocks))
        return doc

class ParserRegistry:
    def __init__(self):
        self._parsers: Dict[str, BaseParser] = {
            "txt": TxtParser(),
            "csv": TxtParser(),
            "pdf": PdfParser(),
            "docx": DocxParser()
        }
        
    def get_parser(self, extension: str) -> BaseParser:
        ext = extension.lower().lstrip(".")
        if ext not in self._parsers:
            raise ValueError(f"No parser available for extension '{ext}'")
        return self._parsers[ext]
